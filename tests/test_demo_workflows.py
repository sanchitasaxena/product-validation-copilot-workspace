import json
import subprocess
import sys
import tempfile
import unittest
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PYTHON = sys.executable
PRD = "examples/research-source-intake/sources/registry-audit-trail-prd.md"
CATALOG = "examples/research-source-intake/sources/target-audience-catalog.synthetic.json"


def run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [PYTHON, *args],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )


class DemoWorkflowTests(unittest.TestCase):
    @staticmethod
    def assert_valid_office_file(path: Path) -> None:
        if shutil.which("officecli") is None:
            raise AssertionError("officecli is required for Office artifact tests.")
        result = subprocess.run(
            ["officecli", "validate", str(path)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode:
            raise AssertionError(result.stdout + result.stderr)

    @staticmethod
    def office_text(path: Path) -> str:
        result = subprocess.run(
            ["officecli", "view", str(path), "text", "--max-lines", "1000"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode:
            raise AssertionError(result.stdout + result.stderr)
        return result.stdout

    def test_qa_source_docs_create_valid_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            result = run(
                "tools/run_qa_bug_bash_workflow.py",
                "--feature", "Registry approval workflow",
                "--product-area", "HCP Terraform Registry",
                "--source-doc", "examples/source-document-intake/sources/feature-brief.md",
                "--source-doc", "examples/source-document-intake/sources/api-notes.txt",
                "--output-dir", directory,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook_path = Path(directory) / "registry-approval-workflow.bug-bash-tracker.xlsx"
            workbook = load_workbook(workbook_path, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Scenarios", "Bug_Tracker", "Summary", "Validation"],
            )
            context = json.loads((Path(directory) / "qa-context.normalized.json").read_text())
            self.assertGreater(len(context["scenarios"]), 0)
            self.assertTrue(all(item["needsEngineeringReview"] for item in context["scenarios"]))
            validation = json.loads((Path(directory) / "validation.json").read_text())
            self.assertEqual(validation["errors"], [])
            self.assert_valid_office_file(workbook_path)

    def test_prd_creates_cuj_mapped_moderated_guide(self):
        with tempfile.TemporaryDirectory() as directory:
            result = run(
                "tools/run_research_workflow.py",
                "--workflow", "testing-guide",
                "--source-doc", PRD,
                "--product-area", "HCP Terraform Registry",
                "--catalog-source", CATALOG,
                "--include-synthetic-catalog",
                "--output-dir", directory,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            context = json.loads((Path(directory) / "research-context.normalized.json").read_text())
            guide = json.loads((Path(directory) / "testing-guide.json").read_text())
            cujs = [item["text"] for item in context["audienceContext"]["cuj"]]
            self.assertEqual([task["title"] for task in guide["tasks"]], cujs)
            self.assertTrue(guide["humanReviewRequired"])
            docx = Path(directory) / "testing-guide.docx"
            self.assertTrue(docx.exists())
            self.assert_valid_office_file(docx)

    def test_prd_creates_discovery_survey_draft(self):
        with tempfile.TemporaryDirectory() as directory:
            result = run(
                "tools/run_research_workflow.py",
                "--workflow", "discovery-survey",
                "--source-doc", PRD,
                "--product-area", "HCP Terraform Registry",
                "--catalog-source", CATALOG,
                "--include-synthetic-catalog",
                "--output-dir", directory,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            survey = json.loads((Path(directory) / "discovery-survey.json").read_text())
            self.assertGreater(len(survey["questions"]), 0)
            self.assertTrue(survey["humanReviewRequired"])
            validation = json.loads((Path(directory) / "discovery-survey.validation.json").read_text())
            self.assertEqual(validation["errors"], [])
            expected = [
                "discovery-survey.docx",
                "discovery-survey.forms-import.docx",
                "discovery-survey.forms.json",
                "discovery-survey.forms.csv",
                "discovery-survey.forms-import.md",
            ]
            for name in expected:
                self.assertTrue((Path(directory) / name).exists(), name)
            self.assert_valid_office_file(Path(directory) / "discovery-survey.docx")
            self.assert_valid_office_file(Path(directory) / "discovery-survey.forms-import.docx")

    def test_monday_request_is_draft_only_and_asks_for_missing_answers(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "monday-request.json"
            result = run(
                "tools/run_research_workflow.py",
                "--workflow", "monday-request",
                "--input", "tests/fixtures/research-context.json",
                "--output", str(output),
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            draft = json.loads(output.read_text())
            self.assertEqual(draft["submissionStatus"], "draft")
            self.assertEqual(draft["fieldMappingStatus"], "unverified-live-form")
            self.assertTrue(draft["humanReviewRequired"])
            self.assertGreater(len(draft["missingQuestions"]), 0)
            self.assertIn(
                "Who is the primary audience for this research?",
                draft["missingQuestions"],
            )
            self.assertTrue(output.with_suffix(".md").exists())
            self.assertTrue(output.with_suffix(".docx").exists())
            self.assert_valid_office_file(output.with_suffix(".docx"))

    def test_remaining_research_workflows_create_valid_docx(self):
        expected_text = {
            "plan": "Decision this research informs",
            "unmoderated-test": "Participant introduction",
            "discovery-guide": "Interview questions",
        }
        with tempfile.TemporaryDirectory() as directory:
            for workflow, heading in expected_text.items():
                output = Path(directory) / f"{workflow}.json"
                result = run(
                    "tools/run_research_workflow.py",
                    "--workflow", workflow,
                    "--input", "tests/fixtures/research-context.json",
                    "--output", str(output),
                )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                docx = output.with_suffix(".docx")
                self.assertTrue(docx.exists())
                self.assert_valid_office_file(docx)
                self.assertIn(heading, self.office_text(docx))

    def test_normal_prose_creates_reviewable_inferred_candidates(self):
        with tempfile.TemporaryDirectory() as directory:
            result = run(
                "tools/build_research_context_from_sources.py",
                "--source-doc", "tests/fixtures/normal-prose-prd.md",
                "--product-area", "Example Registry",
                "--output", str(Path(directory) / "context.json"),
                "--audience-proposal-output", str(Path(directory) / "audience.json"),
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            context = json.loads((Path(directory) / "context.json").read_text())
            audience = context["audienceContext"]
            self.assertEqual(audience["primaryPersona"]["extractionStatus"], "inferred")
            self.assertTrue(audience["jtbd"])
            self.assertTrue(audience["cuj"])
            self.assertTrue(context["requirements"])
            self.assertTrue(context["researchQuestions"])
            claims = [audience["primaryPersona"], *audience["jtbd"], *audience["cuj"]]
            self.assertTrue(all(item["sourceExcerpt"] for item in claims))
            self.assertTrue(all(item["newProposal"] for item in claims))

    def test_synthetic_catalog_is_not_used_without_explicit_opt_in(self):
        with tempfile.TemporaryDirectory() as directory:
            result = run(
                "tools/build_research_context_from_sources.py",
                "--source-doc", PRD,
                "--product-area", "HCP Terraform Registry",
                "--catalog-source", CATALOG,
                "--output", str(Path(directory) / "context.json"),
                "--audience-proposal-output", str(Path(directory) / "audience.json"),
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            context = json.loads((Path(directory) / "context.json").read_text())
            claims = (
                context["audienceContext"]["jtbd"]
                + context["audienceContext"]["cuj"]
            )
            self.assertTrue(all(item["newProposal"] for item in claims))

    def test_missing_source_document_fails_clearly(self):
        result = run(
            "tools/build_research_context_from_sources.py",
            "--source-doc", "does-not-exist.md",
            "--product-area", "Example",
            "--output", "/tmp/unused-context.json",
            "--audience-proposal-output", "/tmp/unused-audience.json",
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("does-not-exist.md", result.stderr)


if __name__ == "__main__":
    unittest.main()
