#!/usr/bin/env python3
"""
Generate an internal QA / bug bash tracker XLSX from a QA context JSON file.

This generator is intentionally schema-first and style-aware:
- input follows schemas/qa-context.schema.json conceptually;
- validation rules come from bug-bash-tracker.rules.json;
- scenario wording should follow qa/scenario-style-guide.md.

Usage:
    python3 generators/xlsx/generate_bug_bash_tracker.py \
        --input generators/xlsx/examples/rauc_sample_qa_context.json \
        --output generators/xlsx/out/RAUC_Bug_Bash_Tracker.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as exc:  # pragma: no cover - explicit operator guidance
    sys.exit(
        "Missing dependency: openpyxl. Install it in your selected Python environment "
        "before running this generator."
    )


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_RULES_PATH = Path(__file__).with_name("bug-bash-tracker.rules.json")

SCENARIO_HEADERS = [
    "ID",
    "Feature",
    "Scenario Description",
    "Prerequisites",
    "Expected Result",
    "Testers",
    "Status",
    "Scenario Type",
    "Source Status",
    "Depends On",
    "Related Bug IDs",
    "Needs Engineering Review",
    "Notes",
]

BUG_HEADERS = [
    "Date Reported",
    "Type",
    "Feature",
    "Scenario ID",
    "Description",
    "Severity",
    "Owner",
    "Evidence Links",
    "Status",
    "Release Blocking",
    "Orphan Classification",
    "Notes",
]

VALIDATION_HEADERS = ["Severity", "Area", "Row ID", "Message"]

STATUS_DISPLAY = {
    "Pending": "⚪ Pending",
    "In Progress": "🟡 In Progress",
    "Passed": "✅ Passed",
    "Failed": "❌ Failed",
    "Blocked": "⛔ Blocked",
}

STATUS_NORMALIZE = {v: k for k, v in STATUS_DISPLAY.items()}
STATUS_NORMALIZE.update({k: k for k in STATUS_DISPLAY})

BUG_STATUS = ["New", "In Progress", "QA", "Resolved", "Deferred"]
BUG_TYPES = ["Bug", "Friction", "Improvement"]
SEVERITIES = ["1-Critical", "2-High", "3-Medium", "4-Low", "Unrated"]

FILL_HEADER = "2D3748"
FILL_LIGHT = "F7F8FA"
FILL_WARN = "FFF3CD"
FILL_FAIL = "FFC7CE"
FILL_PASS = "C6EFCE"
FILL_INFO = "EBF4FA"
WHITE = "FFFFFF"
BORDER_STYLE = "TableStyleMedium2"


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


def normalize_status(status: str | None) -> str:
    if not status:
        return "Pending"
    return STATUS_NORMALIZE.get(status.strip(), status.strip())


def display_status(status: str | None) -> str:
    return STATUS_DISPLAY.get(normalize_status(status), status or "⚪ Pending")


def split_people(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if not value:
        return []
    return [part.strip() for part in str(value).split(",") if part.strip()]


def join_people(value: Any) -> str:
    return ", ".join(split_people(value))


def parse_ids(text: Any) -> list[str]:
    if not text:
        return []
    return sorted(set(re.findall(r"T-\d+", str(text))))


def evidence_links_to_text(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, str):
        return value
    links: list[str] = []
    for item in value:
        if isinstance(item, str):
            links.append(item)
        elif isinstance(item, dict):
            label = item.get("sourceLabel") or item.get("type") or "evidence"
            uri = item.get("uri") or ""
            links.append(f"{label}: {uri}".strip(": "))
    return "\n".join(links)


def scenario_value(scenario: dict[str, Any], key: str, default: Any = "") -> Any:
    aliases = {
        "description": ["description", "scenarioDescription", "Scenario Description"],
        "expectedResult": ["expectedResult", "Expected Result"],
        "prerequisites": ["prerequisites", "Prerequisites"],
    }
    keys = aliases.get(key, [key])
    for candidate in keys:
        if candidate in scenario:
            return scenario[candidate]
    return default


def bug_scenario_ids(bug: dict[str, Any]) -> list[str]:
    ids = []
    if bug.get("scenarioIds"):
        ids.extend(str(item) for item in bug["scenarioIds"])
    if bug.get("scenarioId"):
        ids.append(str(bug["scenarioId"]))
    ids.extend(parse_ids(bug.get("description")))
    return sorted(set(ids))


def validate_context(context: dict[str, Any], rules: dict[str, Any]) -> list[dict[str, str]]:
    warnings: list[dict[str, str]] = []
    scenario_rules = rules.get("scenarioRules", {})
    bug_rules = rules.get("bugRules", {})

    scenarios = context.get("scenarios", [])
    bugs = context.get("bugs", [])
    scenario_ids = [str(s.get("id", "")).strip() for s in scenarios if s.get("id")]
    scenario_id_set = set(scenario_ids)

    for duplicate_id, count in Counter(scenario_ids).items():
        if count > 1:
            warnings.append(rule("error", "Scenarios", duplicate_id, "Duplicate scenario ID."))

    for scenario in scenarios:
        sid = str(scenario.get("id", "")).strip()
        row_label = sid or "(missing ID)"
        if not sid:
            warnings.append(rule("error", "Scenarios", row_label, "Scenario is missing an ID."))
        elif not re.match(scenario_rules.get("idPattern", r"^T-[0-9]+$"), sid):
            warnings.append(rule("error", "Scenarios", row_label, "Scenario ID must use T-### format."))

        for field in scenario_rules.get("requiredFields", []):
            if field == "expectedResult":
                value = scenario_value(scenario, "expectedResult")
            elif field == "description":
                value = scenario_value(scenario, "description")
            else:
                value = scenario.get(field)
            if value in (None, "", []):
                warnings.append(rule("error", "Scenarios", row_label, f"Missing required field: {field}."))

        description = str(scenario_value(scenario, "description", "")).strip()
        for prefix in scenario_rules.get("disallowedDescriptionPrefixes", []):
            if description.lower().startswith(prefix.lower()):
                warnings.append(
                    rule(
                        "warning",
                        "Scenarios",
                        row_label,
                        f"Scenario wording starts with '{prefix}', which reads like research/user-story language.",
                    )
                )

        status = normalize_status(scenario.get("status"))
        if status not in scenario_rules.get("acceptedStatuses", list(STATUS_DISPLAY)):
            warnings.append(rule("warning", "Scenarios", row_label, f"Unknown scenario status: {status}."))

        dependencies = set(scenario.get("dependsOn") or [])
        dependencies.update(parse_ids(scenario_value(scenario, "prerequisites")))
        for dep in sorted(dependencies):
            if dep != sid and dep not in scenario_id_set:
                warnings.append(rule("warning", "Scenarios", row_label, f"Dependency does not exist: {dep}."))

        if scenario.get("sourceStatus") == "inferred" and not scenario.get("needsEngineeringReview", True):
            warnings.append(rule("warning", "Scenarios", row_label, "Inferred scenario should need engineering review."))

    linked_bug_by_scenario: dict[str, list[str]] = defaultdict(list)
    for idx, bug in enumerate(bugs, start=1):
        bid = str(bug.get("id") or f"BUG-{idx:03d}")
        for field in bug_rules.get("requiredFields", []):
            if bug.get(field) in (None, "", []):
                warnings.append(rule("error", "Bug_Tracker", bid, f"Missing required field: {field}."))

        bug_type = bug.get("type")
        if bug_type and bug_type not in bug_rules.get("acceptedTypes", BUG_TYPES):
            warnings.append(rule("warning", "Bug_Tracker", bid, f"Unknown issue type: {bug_type}."))

        severity = bug.get("severity") or "Unrated"
        if severity not in bug_rules.get("acceptedSeverities", SEVERITIES):
            warnings.append(rule("warning", "Bug_Tracker", bid, f"Unknown severity: {severity}."))

        status = bug.get("status") or "New"
        if status not in bug_rules.get("acceptedStatuses", BUG_STATUS):
            warnings.append(rule("warning", "Bug_Tracker", bid, f"Unknown bug status: {status}."))

        scenario_refs = bug_scenario_ids(bug)
        if scenario_refs:
            for sid in scenario_refs:
                linked_bug_by_scenario[sid].append(bid)
                if sid not in scenario_id_set:
                    warnings.append(rule("warning", "Bug_Tracker", bid, f"References missing scenario ID: {sid}."))
        else:
            warnings.append(rule("warning", "Bug_Tracker", bid, "Bug row has no linked scenario ID."))

        unresolved = status in bug_rules.get("unresolvedStatuses", ["New", "In Progress", "QA"])
        high_sev = severity in bug_rules.get("criticalHighSeverities", ["1-Critical", "2-High"])
        if unresolved and high_sev:
            if not bug.get("owner"):
                warnings.append(rule("warning", "Bug_Tracker", bid, "Unresolved critical/high bug has no owner."))
            if not bug.get("evidenceLinks"):
                warnings.append(rule("warning", "Bug_Tracker", bid, "Unresolved critical/high bug has no evidence link."))

    if scenario_rules.get("failedScenarioRequiresBugOrNote", True):
        for scenario in scenarios:
            sid = str(scenario.get("id", "")).strip()
            if normalize_status(scenario.get("status")) == "Failed":
                has_bug = bool(linked_bug_by_scenario.get(sid) or scenario.get("relatedBugIds"))
                has_note = bool(str(scenario.get("notes", "")).strip())
                if not has_bug and not has_note:
                    warnings.append(rule("warning", "Scenarios", sid, "Failed scenario has no linked bug or note."))

    return warnings


def rule(severity: str, area: str, row_id: str, message: str) -> dict[str, str]:
    return {"severity": severity, "area": area, "rowId": row_id, "message": message}


def build_summary(context: dict[str, Any], warnings: list[dict[str, str]]) -> list[list[Any]]:
    scenarios = context.get("scenarios", [])
    bugs = context.get("bugs", [])
    feature_status: dict[str, Counter[str]] = defaultdict(Counter)
    bug_severity = Counter()
    bug_type = Counter()
    owner_load = Counter()
    open_critical_high: list[str] = []
    failed_without_bug: list[str] = []
    orphan_bugs: list[str] = []
    blocked_dependencies: list[str] = []

    for scenario in scenarios:
        feature_status[str(scenario.get("feature", "Unspecified"))][normalize_status(scenario.get("status"))] += 1

    for idx, bug in enumerate(bugs, start=1):
        bid = str(bug.get("id") or f"BUG-{idx:03d}")
        bug_severity[bug.get("severity") or "Unrated"] += 1
        bug_type[bug.get("type") or "Unspecified"] += 1
        if bug.get("owner"):
            owner_load[str(bug["owner"])] += 1
        if bug.get("status") in ("New", "In Progress", "QA") and bug.get("severity") in ("1-Critical", "2-High"):
            open_critical_high.append(bid)
        if not bug_scenario_ids(bug):
            orphan_bugs.append(bid)

    scenario_ids = {str(item.get("id")) for item in scenarios if item.get("id")}
    linked_ids = {sid for bug in bugs for sid in bug_scenario_ids(bug)}
    for scenario in scenarios:
        sid = str(scenario.get("id", ""))
        if normalize_status(scenario.get("status")) == "Failed" and sid not in linked_ids and not scenario.get("notes"):
            failed_without_bug.append(sid)
        if normalize_status(scenario.get("status")) == "Blocked":
            missing = [dep for dep in scenario.get("dependsOn", []) if dep not in scenario_ids]
            if missing:
                blocked_dependencies.append(f"{sid}: {', '.join(missing)}")

    rows: list[list[Any]] = [
        ["Summary Item", "Value"],
        ["Feature", context.get("feature", "")],
        ["Product Area", context.get("productArea", "")],
        ["QA Goal", context.get("qaGoal", "")],
        ["Scenario Count", len(scenarios)],
        ["Bug/Issue Count", len(bugs)],
        ["Validation Errors", sum(1 for w in warnings if w["severity"] == "error")],
        ["Validation Warnings", sum(1 for w in warnings if w["severity"] == "warning")],
        ["Open Critical/High Bugs", ", ".join(open_critical_high) or "None"],
        ["Failed Scenarios Without Bugs", ", ".join(failed_without_bug) or "None"],
        ["Orphan/Exploratory Bugs", ", ".join(orphan_bugs) or "None"],
        ["Blocked Dependency Chains", "; ".join(blocked_dependencies) or "None"],
        [],
        ["Scenario Status by Feature", ""],
    ]

    for feature, counts in sorted(feature_status.items()):
        total = sum(counts.values())
        status_bits = ", ".join(f"{status}: {counts.get(status, 0)}" for status in STATUS_DISPLAY)
        rows.append([feature, f"{status_bits}; Total: {total}"])

    rows.extend([[], ["Bugs by Severity", ""]])
    for severity, count in sorted(bug_severity.items()):
        rows.append([severity, count])

    rows.extend([[], ["Bugs by Type", ""]])
    for issue_type, count in sorted(bug_type.items()):
        rows.append([issue_type, count])

    rows.extend([[], ["Owner Workload", ""]])
    for owner, count in sorted(owner_load.items()):
        rows.append([owner, count])

    return rows


def set_title(ws: Any, title: str, column_count: int) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos", bold=True, size=18, color="1F1F1F")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    ws.row_dimensions[1].height = 30


def style_header(row: Any) -> None:
    for cell in row:
        cell.font = Font(name="Aptos", bold=True, size=11, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=FILL_HEADER)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_body(ws: Any) -> None:
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.font = Font(name="Aptos", size=11, color="1F1F1F")
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def add_table(ws: Any, name: str, start_row: int, end_row: int, end_col: int) -> None:
    if end_row < start_row:
        return
    ref = f"A{start_row}:{openpyxl.utils.get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=BORDER_STYLE,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_list_validation(ws: Any, cell_range: str, values: list[str]) -> None:
    formula = '"' + ",".join(values) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.sqref = cell_range
    ws.add_data_validation(dv)


def add_status_formatting(ws: Any, cell_range: str) -> None:
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f'ISNUMBER(SEARCH("Passed",G3))'], fill=PatternFill("solid", fgColor=FILL_PASS)))
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f'ISNUMBER(SEARCH("Failed",G3))'], fill=PatternFill("solid", fgColor=FILL_FAIL)))
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f'ISNUMBER(SEARCH("Blocked",G3))'], fill=PatternFill("solid", fgColor=FILL_WARN)))
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f'ISNUMBER(SEARCH("In Progress",G3))'], fill=PatternFill("solid", fgColor=FILL_WARN)))


def build_workbook(context: dict[str, Any], warnings: list[dict[str, str]]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    build_scenarios_sheet(wb, context)
    build_bug_sheet(wb, context)
    build_summary_sheet(wb, context, warnings)
    build_validation_sheet(wb, warnings)

    return wb


def build_scenarios_sheet(wb: Workbook, context: dict[str, Any]) -> None:
    ws = wb.create_sheet("Scenarios")
    title = f"{context.get('feature', 'Feature')} - Scenarios"
    set_title(ws, title, len(SCENARIO_HEADERS))
    ws.append([])
    ws.append(SCENARIO_HEADERS)
    style_header(ws[3])

    for scenario in context.get("scenarios", []):
        depends_on = scenario.get("dependsOn") or parse_ids(scenario_value(scenario, "prerequisites"))
        row = [
            scenario.get("id", ""),
            scenario.get("feature", ""),
            scenario_value(scenario, "description", ""),
            scenario_value(scenario, "prerequisites", ""),
            scenario_value(scenario, "expectedResult", ""),
            join_people(scenario.get("testers")),
            display_status(scenario.get("status")),
            scenario.get("scenarioType", ""),
            scenario.get("sourceStatus", ""),
            ", ".join(depends_on),
            ", ".join(scenario.get("relatedBugIds", [])),
            "Yes" if scenario.get("needsEngineeringReview") else "No",
            scenario.get("notes", ""),
        ]
        ws.append(row)

    style_body(ws)
    end_row = max(ws.max_row, 3)
    add_table(ws, "Scenarios", 3, end_row, len(SCENARIO_HEADERS))
    add_list_validation(ws, f"G4:G{max(end_row, 200)}", list(STATUS_DISPLAY.values()))
    add_list_validation(ws, f"H4:H{max(end_row, 200)}", ["UI", "API", "CLI", "Permissions", "Data", "Regression", "Edge Case", "Setup", "Exploratory"])
    add_list_validation(ws, f"I4:I{max(end_row, 200)}", ["source-backed", "inferred", "regression", "edge-case", "exploratory"])
    add_list_validation(ws, f"L4:L{max(end_row, 200)}", ["Yes", "No"])
    add_status_formatting(ws, f"G4:G{max(end_row, 200)}")
    set_widths(ws, [14, 28, 52, 34, 62, 26, 18, 20, 18, 18, 20, 22, 36])


def build_bug_sheet(wb: Workbook, context: dict[str, Any]) -> None:
    ws = wb.create_sheet("Bug_Tracker")
    title = f"{context.get('feature', 'Feature')} - Bug Tracker"
    set_title(ws, title, len(BUG_HEADERS))
    ws.append([])
    ws.append(BUG_HEADERS)
    style_header(ws[3])

    for idx, bug in enumerate(context.get("bugs", []), start=1):
        scenario_ids = bug_scenario_ids(bug)
        row = [
            bug.get("dateReported", date.today().isoformat()),
            bug.get("type", ""),
            bug.get("feature", ""),
            ", ".join(scenario_ids),
            bug.get("description", ""),
            bug.get("severity", "Unrated"),
            bug.get("owner", ""),
            evidence_links_to_text(bug.get("evidenceLinks")),
            bug.get("status", "New"),
            "Yes" if bug.get("releaseBlocking") else "No",
            bug.get("orphanClassification") or ("not-orphan" if scenario_ids else "orphan"),
            bug.get("notes", ""),
        ]
        ws.append(row)

    style_body(ws)
    end_row = max(ws.max_row, 3)
    add_table(ws, "BugTable", 3, end_row, len(BUG_HEADERS))
    add_list_validation(ws, f"B4:B{max(end_row, 200)}", BUG_TYPES)
    add_list_validation(ws, f"F4:F{max(end_row, 200)}", SEVERITIES)
    add_list_validation(ws, f"I4:I{max(end_row, 200)}", BUG_STATUS)
    add_list_validation(ws, f"J4:J{max(end_row, 200)}", ["Yes", "No"])
    add_list_validation(ws, f"K4:K{max(end_row, 200)}", ["orphan", "exploratory", "no-scenario", "not-orphan"])
    set_widths(ws, [18, 16, 28, 18, 68, 16, 28, 52, 16, 18, 22, 36])


def build_summary_sheet(wb: Workbook, context: dict[str, Any], warnings: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("Summary")
    rows = build_summary(context, warnings)
    set_title(ws, f"{context.get('feature', 'Feature')} - Release Readiness Summary", 2)
    ws.append([])
    for row in rows:
        ws.append(row)
    style_header(ws[3])
    style_body(ws)
    set_widths(ws, [38, 96])


def build_validation_sheet(wb: Workbook, warnings: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("Validation")
    set_title(ws, "Validation Warnings", len(VALIDATION_HEADERS))
    ws.append([])
    ws.append(VALIDATION_HEADERS)
    style_header(ws[3])

    if warnings:
        for warning in warnings:
            ws.append([warning["severity"], warning["area"], warning["rowId"], warning["message"]])
    else:
        ws.append(["info", "Workbook", "", "No validation issues found."])

    style_body(ws)
    add_table(ws, "ValidationTable", 3, ws.max_row, len(VALIDATION_HEADERS))
    set_widths(ws, [14, 18, 18, 96])
    for row in ws.iter_rows(min_row=4, max_col=1):
        cell = row[0]
        if cell.value == "error":
            cell.fill = PatternFill("solid", fgColor=FILL_FAIL)
        elif cell.value == "warning":
            cell.fill = PatternFill("solid", fgColor=FILL_WARN)
        else:
            cell.fill = PatternFill("solid", fgColor=FILL_INFO)


def set_widths(ws: Any, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    ws.freeze_panes = "A4"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate a QA bug bash tracker XLSX.")
    parser.add_argument("--input", required=True, help="Path to QA context JSON.")
    parser.add_argument("--output", required=True, help="Path to output XLSX.")
    parser.add_argument("--rules", default=str(DEFAULT_RULES_PATH), help="Path to bug bash tracker rules JSON.")
    parser.add_argument("--validation-report", help="Optional path to write validation warnings JSON.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    context_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    rules_path = Path(args.rules).expanduser().resolve()

    context = load_json(context_path)
    rules = load_json(rules_path)
    warnings = validate_context(context, rules)

    wb = build_workbook(context, warnings)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    if args.validation_report:
        write_json(
            Path(args.validation_report).expanduser().resolve(),
            {
                "errors": [item for item in warnings if item["severity"] == "error"],
                "warnings": [item for item in warnings if item["severity"] == "warning"],
                "issues": warnings,
            },
        )

    errors = sum(1 for warning in warnings if warning["severity"] == "error")
    warning_count = sum(1 for warning in warnings if warning["severity"] == "warning")
    print(f"Done: {output_path}")
    print(f"Validation: {errors} error(s), {warning_count} warning(s)")
    if errors:
        sys.exit(1)


if __name__ == "__main__":
    main()
