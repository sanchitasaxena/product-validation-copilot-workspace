#!/usr/bin/env python3
"""
Import an existing internal QA / bug bash tracker XLSX into QA context JSON.

The importer is designed for the original RAUC tracker shape:
- `Scenarios` sheet with ID, Feature, Scenario Description, Prerequisites,
  Expected Result, Testers, Status.
- `Bug_Tracker` sheet with Date Reported, Type, Feature, Description,
  Severity, Owner, Link to screenshot, Status, optional notes/category.

It preserves source wording and messy real-world values while normalizing enough
structure for `generate_bug_bash_tracker.py`.

Usage:
    python3 generators/xlsx/import_bug_bash_tracker.py \
        --input "../RAUC Example/rauc context/RAUC_Bug_Tracker_2026_copy.xlsx" \
        --output generators/xlsx/out/RAUC_Bug_Tracker_2026.qa-context.json
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:  # pragma: no cover - explicit operator guidance
    sys.exit(
        "Missing dependency: openpyxl. Install it in your selected Python environment "
        "before running this importer."
    )


DEFAULT_FEATURE = "Registry Artifact Usage Control (RAUC)"
DEFAULT_PRODUCT_AREA = "HCP Terraform Registry"
DEFAULT_SCENARIOS_SHEET = "Scenarios"
DEFAULT_BUGS_SHEET = "Bug_Tracker"
DEFAULT_STYLE_SOURCE = "Original RAUC bug tracker engineering-authored scenarios"

STATUS_ALIASES = {
    "⚪ Pending": "Pending",
    "🟡 In Progress": "In Progress",
    "✅ Passed": "Passed",
    "❌ Failed": "Failed",
    "pending": "Pending",
    "in progress": "In Progress",
    "passed": "Passed",
    "failed": "Failed",
    "blocked": "Blocked",
}

BUG_STATUS_ALIASES = {
    "new": "New",
    "in progress": "In Progress",
    "qa": "QA",
    "resolved": "Resolved",
    "deferred": "Deferred",
}

ISSUE_TYPES = {"Bug", "Friction", "Improvement"}
SEVERITIES = {"1-Critical", "2-High", "3-Medium", "4-Low", "Unrated"}
URL_RE = re.compile(r"https?://\S+", re.IGNORECASE)
SCENARIO_ID_RE = re.compile(r"\bT-?\s*(\d{3})\b", re.IGNORECASE)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def date_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return cell_text(value)


def normalize_status(value: Any) -> str:
    text = cell_text(value)
    if not text:
        return "Pending"
    return STATUS_ALIASES.get(text, STATUS_ALIASES.get(text.lower(), text))


def normalize_bug_status(value: Any) -> str:
    text = cell_text(value)
    if not text:
        return "New"
    return BUG_STATUS_ALIASES.get(text, BUG_STATUS_ALIASES.get(text.lower(), text))


def normalize_issue_type(value: Any) -> tuple[str, list[str]]:
    text = cell_text(value)
    if not text:
        return "Bug", ["Missing issue type in source tracker; defaulted to Bug for schema compatibility."]
    if text in ISSUE_TYPES:
        return text, []
    return text, [f"Unknown issue type in source tracker: {text}."]


def normalize_severity(value: Any) -> tuple[str, list[str]]:
    text = cell_text(value)
    if not text:
        return "Unrated", ["Missing severity in source tracker; defaulted to Unrated."]
    if text in SEVERITIES:
        return text, []
    return text, [f"Unknown severity in source tracker: {text}."]


def split_people(value: Any) -> list[str]:
    text = cell_text(value)
    if not text:
        return []
    return [part.strip() for part in text.split(",") if part.strip()]


def normalize_scenario_id(value: Any) -> str:
    text = cell_text(value)
    match = SCENARIO_ID_RE.search(text)
    if match:
        return f"T-{match.group(1)}"
    return text


def parse_scenario_ids(*values: Any) -> list[str]:
    ids: set[str] = set()
    for value in values:
        for match in SCENARIO_ID_RE.finditer(cell_text(value)):
            ids.add(f"T-{match.group(1)}")
    return sorted(ids)


def infer_formula_scenario_id(row_number: int, header_row: int, formula_value: Any) -> str:
    text = cell_text(formula_value)
    if not text.startswith("="):
        return normalize_scenario_id(text)
    if '"T-"' in text and "ROW()" in text:
        return f"T-{100 + row_number - header_row:03d}"
    return ""


def infer_scenario_type(feature: str, description: str, expected_result: str) -> str:
    haystack = f"{feature} {description} {expected_result}".lower()
    if "terraform cli" in haystack or " cli" in haystack:
        return "CLI"
    if "api" in haystack:
        return "API"
    if "permission" in haystack or "premium tier" in haystack or "usage control" in haystack:
        return "Permissions"
    if "setup" in haystack or "configured" in haystack:
        return "Setup"
    if "regression" in haystack:
        return "Regression"
    if "edge" in haystack or "maximum" in haystack or "large number" in haystack:
        return "Edge Case"
    return "UI"


def evidence_type(uri: str) -> str:
    lowered = uri.lower()
    if "slack.com" in lowered:
        return "slack"
    if "box.com" in lowered:
        return "box"
    if any(lowered.endswith(ext) for ext in (".png", ".jpg", ".jpeg", ".gif", ".webp")):
        return "screenshot"
    if "github.com" in lowered:
        return "github"
    return "other"


def make_evidence_links(*values: Any) -> list[dict[str, str]]:
    links: list[dict[str, str]] = []
    seen: set[str] = set()
    for value in values:
        for uri in URL_RE.findall(cell_text(value)):
            if uri in seen:
                continue
            seen.add(uri)
            links.append(
                {
                    "id": f"E-{len(links) + 1:03d}",
                    "type": evidence_type(uri),
                    "uri": uri,
                    "sourceLabel": "Imported evidence link",
                }
            )
    return links


def find_header_row(ws: Worksheet, required_headers: set[str]) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=False):
        values = {cell_text(cell.value) for cell in row}
        if required_headers.issubset(values):
            return row[0].row
    raise ValueError(f"Could not find required headers in sheet '{ws.title}'.")


def header_map(ws: Worksheet, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for cell in ws[header_row]:
        header = cell_text(cell.value)
        if header:
            mapping[header] = cell.column
    return mapping


def sheet_table_end_row(ws: Worksheet, fallback_header_row: int) -> int:
    if ws.tables:
        end_rows: list[int] = []
        for table in ws.tables.values():
            ref = getattr(table, "ref", "")
            if ":" not in ref:
                continue
            end_cell = ref.split(":")[-1]
            match = re.search(r"(\d+)$", end_cell)
            if match:
                end_rows.append(int(match.group(1)))
        if end_rows:
            return max(end_rows)
    return ws.max_row or fallback_header_row


def value_at(ws: Worksheet, row: int, headers: dict[str, int], header: str) -> Any:
    column = headers.get(header)
    if not column:
        return None
    return ws.cell(row=row, column=column).value


def import_scenarios(wb: Any, wb_values: Any, sheet_name: str) -> tuple[list[dict[str, Any]], set[str]]:
    ws = wb[sheet_name]
    ws_values = wb_values[sheet_name]
    header_row = find_header_row(ws, {"ID", "Feature", "Scenario Description"})
    headers = header_map(ws, header_row)
    end_row = sheet_table_end_row(ws, header_row)
    scenarios: list[dict[str, Any]] = []
    scenario_ids: set[str] = set()

    for row in range(header_row + 1, end_row + 1):
        description = cell_text(value_at(ws, row, headers, "Scenario Description"))
        feature = cell_text(value_at(ws, row, headers, "Feature"))
        if not description and not feature:
            continue

        cached_id = value_at(ws_values, row, headers, "ID")
        formula_or_id = value_at(ws, row, headers, "ID")
        scenario_id = normalize_scenario_id(cached_id) or infer_formula_scenario_id(row, header_row, formula_or_id)
        if not scenario_id:
            scenario_id = f"T-{100 + row - header_row:03d}"

        prerequisites = cell_text(value_at(ws, row, headers, "Prerequisites"))
        expected_result = cell_text(value_at(ws, row, headers, "Expected Result"))
        testers = split_people(value_at(ws, row, headers, "Testers"))
        status = normalize_status(value_at(ws, row, headers, "Status"))
        depends_on = [dep for dep in parse_scenario_ids(prerequisites) if dep != scenario_id]
        scenario_type = infer_scenario_type(feature, description, expected_result)

        notes: list[str] = []
        if not expected_result:
            notes.append("Imported from original tracker with missing expected result.")
        if not testers:
            notes.append("Imported from original tracker with no tester assigned.")
        if not cell_text(value_at(ws, row, headers, "Status")):
            notes.append("Imported from original tracker with no status; defaulted to Pending.")

        scenarios.append(
            {
                "id": scenario_id,
                "feature": feature or "Unspecified",
                "description": description,
                "prerequisites": prerequisites,
                "expectedResult": expected_result,
                "testers": testers,
                "status": status,
                "scenarioType": scenario_type,
                "sourceStatus": "source-backed",
                "dependsOn": depends_on,
                "relatedBugIds": [],
                "sourceRefs": [f"{sheet_name}!{row}:{row}"],
                "needsEngineeringReview": False,
                "styleWarnings": [],
                "notes": " ".join(notes),
            }
        )
        scenario_ids.add(scenario_id)

    return scenarios, scenario_ids


def import_bugs(ws: Worksheet, sheet_name: str, scenario_ids: set[str]) -> tuple[list[dict[str, Any]], list[str]]:
    header_row = find_header_row(ws, {"Date Reported", "Type", "Feature", "Description"})
    headers = header_map(ws, header_row)
    end_row = sheet_table_end_row(ws, header_row)
    bugs: list[dict[str, Any]] = []
    skipped_rows: list[str] = []

    for row in range(header_row + 1, end_row + 1):
        description = cell_text(value_at(ws, row, headers, "Description"))
        feature = cell_text(value_at(ws, row, headers, "Feature"))
        issue_type_raw = value_at(ws, row, headers, "Type")
        severity_raw = value_at(ws, row, headers, "Severity")
        owner_raw = cell_text(value_at(ws, row, headers, "Owner"))
        evidence_raw = cell_text(value_at(ws, row, headers, "Link to screenshot"))
        status_raw = value_at(ws, row, headers, "Status")
        category = cell_text(value_at(ws, row, headers, "Column1"))

        if not any([description, feature, issue_type_raw, severity_raw, owner_raw, evidence_raw, status_raw, category]):
            continue
        if not description:
            skipped_rows.append(f"{sheet_name}!{row}:{row} skipped because it has no description.")
            continue

        issue_type, warnings = normalize_issue_type(issue_type_raw)
        severity, severity_warnings = normalize_severity(severity_raw)
        warnings.extend(severity_warnings)
        status = normalize_bug_status(status_raw)

        owner = owner_raw
        evidence_values = [evidence_raw]
        if URL_RE.fullmatch(owner_raw):
            evidence_values.append(owner_raw)
            owner = ""
            warnings.append("Owner column contained a URL in source tracker; moved it to evidenceLinks.")

        evidence_links = make_evidence_links(*evidence_values)
        scenario_refs = parse_scenario_ids(description)
        unknown_scenario_refs = [sid for sid in scenario_refs if sid not in scenario_ids]
        if unknown_scenario_refs:
            warnings.append(f"References scenario IDs not found in Scenarios sheet: {', '.join(unknown_scenario_refs)}.")

        if scenario_refs:
            orphan_classification = "not-orphan"
        elif description.lower().startswith("no scenario"):
            orphan_classification = "no-scenario"
        else:
            orphan_classification = "orphan"

        notes = []
        if category:
            notes.append(f"Original category/notes: {category}")
        notes.extend(warnings)

        bug = {
            "id": f"BUG-{len(bugs) + 1:03d}",
            "dateReported": date_text(value_at(ws, row, headers, "Date Reported")),
            "type": issue_type,
            "feature": feature or "Unspecified",
            "scenarioId": scenario_refs[0] if scenario_refs else "",
            "scenarioIds": scenario_refs,
            "description": description,
            "severity": severity,
            "owner": owner,
            "evidenceLinks": evidence_links,
            "status": status,
            "releaseBlocking": status in {"New", "In Progress", "QA"} and severity in {"1-Critical", "2-High"},
            "orphanClassification": orphan_classification,
            "validationWarnings": warnings,
            "notes": " ".join(notes),
            "sourceRefs": [f"{sheet_name}!{row}:{row}"],
        }
        bugs.append(bug)

    return bugs, skipped_rows


def link_bugs_to_scenarios(scenarios: list[dict[str, Any]], bugs: list[dict[str, Any]]) -> None:
    by_scenario: dict[str, list[str]] = {scenario["id"]: [] for scenario in scenarios}
    for bug in bugs:
        for scenario_id in bug.get("scenarioIds", []):
            if scenario_id in by_scenario:
                by_scenario[scenario_id].append(bug["id"])
    for scenario in scenarios:
        scenario["relatedBugIds"] = by_scenario.get(scenario["id"], [])


def build_context(args: argparse.Namespace) -> dict[str, Any]:
    input_path = Path(args.input).expanduser().resolve()
    wb = load_workbook(input_path, data_only=False)
    wb_values = load_workbook(input_path, data_only=True)

    if args.scenarios_sheet not in wb.sheetnames:
        raise ValueError(f"Workbook does not contain scenarios sheet '{args.scenarios_sheet}'.")
    if args.bugs_sheet not in wb.sheetnames:
        raise ValueError(f"Workbook does not contain bug sheet '{args.bugs_sheet}'.")

    scenarios, scenario_ids = import_scenarios(wb, wb_values, args.scenarios_sheet)
    bugs, skipped_bug_rows = import_bugs(wb[args.bugs_sheet], args.bugs_sheet, scenario_ids)
    link_bugs_to_scenarios(scenarios, bugs)

    feature_areas = sorted(
        {scenario["feature"] for scenario in scenarios if scenario.get("feature")}
        | {bug["feature"] for bug in bugs if bug.get("feature") and bug.get("feature") != "Unspecified"}
    )

    return {
        "feature": args.feature,
        "productArea": args.product_area,
        "qaGoal": args.qa_goal,
        "scenarioStyleSource": DEFAULT_STYLE_SOURCE,
        "styleGuideRef": "qa/scenario-style-guide.md",
        "validationRulesRef": "generators/xlsx/bug-bash-tracker.rules.json",
        "sourceDocs": [str(input_path)],
        "featureAreas": feature_areas,
        "audienceGroundingRequired": False,
        "scenarios": scenarios,
        "bugs": bugs,
        "importMetadata": {
            "sourceWorkbook": str(input_path),
            "scenariosSheet": args.scenarios_sheet,
            "bugsSheet": args.bugs_sheet,
            "scenarioCount": len(scenarios),
            "bugCount": len(bugs),
            "skippedRows": skipped_bug_rows,
            "notes": [
                "Formula-based scenario IDs were resolved from cached values or inferred from source row position.",
                "Source wording was preserved; missing fields were annotated rather than rewritten.",
                "URLs found in the owner column were treated as evidence links to handle legacy column drift.",
            ],
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import an existing bug bash XLSX tracker into QA context JSON.")
    parser.add_argument("--input", required=True, help="Path to existing tracker XLSX.")
    parser.add_argument("--output", required=True, help="Path to output QA context JSON.")
    parser.add_argument("--feature", default=DEFAULT_FEATURE, help="Feature name for the QA context.")
    parser.add_argument("--product-area", default=DEFAULT_PRODUCT_AREA, help="Product area for the QA context.")
    parser.add_argument("--qa-goal", default="bug-bash", choices=["bug-bash", "release-readiness", "regression", "exploratory"])
    parser.add_argument("--scenarios-sheet", default=DEFAULT_SCENARIOS_SHEET, help="Name of scenarios sheet.")
    parser.add_argument("--bugs-sheet", default=DEFAULT_BUGS_SHEET, help="Name of bug tracker sheet.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    context = build_context(args)
    output_path = Path(args.output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(context, indent=2) + "\n", encoding="utf-8")
    print(f"Done: {output_path}")
    print(f"Imported: {len(context['scenarios'])} scenario(s), {len(context['bugs'])} bug/issue row(s)")


if __name__ == "__main__":
    main()
